"""Share of search: where IMECO ranks when a shopper searches a marketplace.

Once a day the app types each keyword into Amazon, Flipkart and Blinkit and records the first
page of results in order — every listing, ours or not, organic or paid. The dashboard's Search
tab turns that into "IMECO's best rank for 'kitchen tissue roll' on Flipkart is #11", "Origami
holds 4 of the top 10", and "who is buying ads on our keywords".

    python -m desktop.search                 # search every keyword now, print, store nothing
    python -m desktop.search "toilet roll"   # just one

WHICH RESULT IS WHOSE. Three rules, most certain first:
  1. Its listing id (ASIN / Flipkart pid / Blinkit prid) is one of the sheet's links — then it
     IS that product, IMECO's or a named competitor's, and nothing is guessed.
  2. Blinkit says the brand outright in its search data; that is taken as given.
  3. The title names a brand we know — IMECO's own (see OWN_BRANDS) or one from the Competitors
     sheet. Everything else is left unattributed and counted as "other" — inventing a brand from
     the first word of a title would turn "Kitchen Tissue Paper Roll" into a brand called Kitchen.
"""

from __future__ import annotations

import html as htmllib
import json
import re
import sys
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from urllib.parse import quote, quote_plus

from openpyxl import load_workbook

from .catalogue import LISTING_ID, active_workbook
from .fetch import Fetcher
from .models import MARKETPLACES, Product

#: Brand names IMECO sells under. Some of the catalogue is branded "Good Habbit" (and is spelt
#: "Good Habit" on at least one listing), so a title naming either is ours.
OWN_BRANDS = ("IMECO", "Good Habbit", "Good Habit")

#: Searched when the workbook has no "Keywords" sheet: what a shopper would type for each of
#: the catalogue's product types. A "Keywords" sheet (one "Keyword" column, optional
#: "Marketplaces" column) replaces this list entirely.
DEFAULT_KEYWORDS = (
    "kitchen tissue roll",
    "kitchen towel roll",
    "toilet roll",
    "facial tissue",
    "facial tissue box",
    "pocket tissue",
    "car tissue",
    "paper napkins",
    "garbage bags",
    "wet wipes",
    "baby wipes",
    "disposable plates",
    "wooden cutlery",
    "paper cups",
    "bamboo toothbrush",
    "tongue cleaner",
    "neem comb",
    "hair brush",
    "bath loofah",
)

KEYWORD_SHEET = "keywords"

#: The hour of day from which the daily search is due. Ranks are read once a day, as
#: 42signals does; earlier than this the stores are still settling the day's ads.
SEARCH_FROM_HOUR = 9


@dataclass
class SearchResult:
    keyword: str
    marketplace: str
    position: int
    """1-based, in the order the marketplace showed it, ads included — what a shopper sees."""
    listing_id: str
    title: str
    price: float | None
    sponsored: bool
    brand: str | None = None
    """Attributed brand, "IMECO" for ours. None when it could not be told — see the module doc."""
    matched_sku: str | None = None
    """The sheet product this IS, when its listing id is one of the sheet's links."""


# ── where to search ─────────────────────────────────────────────────────────────────────


def search_request(marketplace: str, keyword: str) -> tuple[str, dict | None]:
    """(url, json body or None). Blinkit's search is a JSON API taking a POST; the other two
    are ordinary result pages."""
    if marketplace == "amazon":
        return f"https://www.amazon.in/s?k={quote_plus(keyword)}", None
    if marketplace == "flipkart":
        return f"https://www.flipkart.com/search?q={quote_plus(keyword)}", None
    if marketplace == "blinkit":
        return (
            f"https://blinkit.com/v1/layout/search?q={quote(keyword)}&search_type=type_to_search",
            {},
        )
    raise ValueError(marketplace)


# ── reading result pages ────────────────────────────────────────────────────────────────
# Each returns results without keyword/marketplace attribution; `search()` fills those in.

_AMAZON_BLOCK = re.compile(r'(?=<div[^>]*data-component-type="s-search-result")')
_AMAZON_ASIN = re.compile(r'data-asin="([A-Z0-9]{10})"')
_AMAZON_TITLE = re.compile(r'<h2[^>]*aria-label="([^"]+)"')
_AMAZON_TITLE_SPAN = re.compile(r"<h2[^>]*>\s*(?:<a[^>]*>)?\s*<span[^>]*>([^<]+)</span>", re.S)
_AMAZON_PRICE = re.compile(r'class="a-price-whole">\s*([\d,]+)')
# Amazon labels a paid slot both in the title ("Sponsored Ad - ...") and with a badge class.
_AMAZON_SPONSORED = re.compile(r'puis-sponsored-label|AdHolder|aria-label="Sponsored Ad - ')


def _price(raw: str | None) -> float | None:
    if not raw:
        return None
    try:
        value = float(str(raw).replace(",", ""))
    except ValueError:
        return None
    return value if value > 0 else None


def parse_amazon(page: str) -> list[SearchResult]:
    out: list[SearchResult] = []
    for block in _AMAZON_BLOCK.split(page)[1:]:
        asin = _AMAZON_ASIN.search(block[:2000])
        if not asin:
            continue
        title = _AMAZON_TITLE.search(block) or _AMAZON_TITLE_SPAN.search(block)
        text = htmllib.unescape(title.group(1)).strip() if title else ""
        sponsored = bool(_AMAZON_SPONSORED.search(block)) or text.startswith("Sponsored Ad - ")
        price = _AMAZON_PRICE.search(block)
        out.append(
            SearchResult(
                keyword="",
                marketplace="amazon",
                position=len(out) + 1,
                listing_id=asin.group(1),
                title=text.removeprefix("Sponsored Ad - "),
                price=_price(price.group(1) if price else None),
                sponsored=sponsored,
            )
        )
    return out


_FLIPKART_ID = re.compile(r'data-id="([A-Z0-9]{16})"')
_FLIPKART_TITLE = re.compile(r'(?:title|alt)="([^"]{8,})"')
_FLIPKART_PRICE = re.compile(r">₹([\d,]+)<")
# A paid slot carries a real banner id; organic results carry "-1".
_FLIPKART_AD = re.compile(r'"adInfo":\{"bannerId":"(\d+)"[^{}]*?"productId":"([A-Z0-9]{16})"')


def parse_flipkart(page: str) -> list[SearchResult]:
    ads = {pid for _, pid in _FLIPKART_AD.findall(page)}
    starts: dict[str, int] = {}
    for m in _FLIPKART_ID.finditer(page):
        starts.setdefault(m.group(1), m.start())
    ordered = sorted(starts.items(), key=lambda kv: kv[1])

    out: list[SearchResult] = []
    for i, (pid, at) in enumerate(ordered):
        end = ordered[i + 1][1] if i + 1 < len(ordered) else at + 8000
        card = page[at:end]
        title = _FLIPKART_TITLE.search(card)
        price = _FLIPKART_PRICE.search(card)
        out.append(
            SearchResult(
                keyword="",
                marketplace="flipkart",
                position=len(out) + 1,
                listing_id=pid,
                title=htmllib.unescape(title.group(1)).strip() if title else "",
                price=_price(price.group(1) if price else None),
                sponsored=pid in ads,
            )
        )
    return out


def parse_blinkit(text: str) -> list[SearchResult]:
    data = json.loads(text)
    out: list[SearchResult] = []
    for snippet in (data.get("response") or {}).get("snippets") or []:
        attrs = (snippet.get("tracking") or {}).get("common_attributes") or {}
        pid = attrs.get("product_id")
        if not pid:
            continue  # headers, banners, "showing results for"
        body = snippet.get("data") or {}
        item = ((body.get("atc_action") or {}).get("add_to_cart") or {}).get("cart_item") or {}
        name = (body.get("name") or {}).get("text") or item.get("product_name") or ""
        out.append(
            SearchResult(
                keyword="",
                marketplace="blinkit",
                position=len(out) + 1,
                listing_id=str(pid),
                title=str(name).strip(),
                price=_price(item.get("price")),
                sponsored=attrs.get("badge") == "AD" or bool(attrs.get("ads_type")),
                brand=(str(attrs["brand"]).strip() or None) if attrs.get("brand") else None,
            )
        )
    return out


PARSERS = {"amazon": parse_amazon, "flipkart": parse_flipkart, "blinkit": parse_blinkit}


# ── whose is it ─────────────────────────────────────────────────────────────────────────


def attribute(results: list[SearchResult], products: list[Product]) -> None:
    """Fill `matched_sku` and `brand` in place, by the three rules in the module doc."""
    by_listing: dict[tuple[str, str], Product] = {}
    for p in products:
        for marketplace, url in p.urls.items():
            m = LISTING_ID[marketplace].search(url)
            if m:
                by_listing[(marketplace, m.group(1))] = p

    known = {b.lower(): "IMECO" for b in OWN_BRANDS}
    for p in products:
        if p.competitor_of:
            known.setdefault(p.brand.lower(), p.brand)
    # Longest first, so "Good Habbit" wins over a shorter brand that happens to be inside it.
    patterns = [
        (re.compile(rf"(?<![\w]){re.escape(name)}(?![\w])", re.I), brand)
        for name, brand in sorted(known.items(), key=lambda kv: -len(kv[0]))
    ]

    for r in results:
        product = by_listing.get((r.marketplace, r.listing_id))
        if product:
            r.matched_sku = product.sku
            r.brand = "IMECO" if product.competitor_of is None else product.brand
            continue
        if r.brand:
            r.brand = known.get(r.brand.lower(), r.brand)
            continue
        r.brand = next((brand for pattern, brand in patterns if pattern.search(r.title)), None)


# ── what to search, and when ────────────────────────────────────────────────────────────


def load_keywords(path: Path | None = None) -> list[tuple[str, tuple[str, ...]]]:
    """[(keyword, marketplaces)] from the workbook's "Keywords" sheet, else the defaults.

    A blank "Marketplaces" cell means all three. Duplicates are dropped, case-insensitively.
    """
    path = path or active_workbook()
    sheet = None
    if path.exists():
        book = load_workbook(path, data_only=True, read_only=True)
        sheet = next((ws for ws in book.worksheets if ws.title.strip().lower() == KEYWORD_SHEET), None)

    if sheet is None:
        return [(k, MARKETPLACES) for k in DEFAULT_KEYWORDS]

    rows = sheet.iter_rows(values_only=True)
    header = [str(c or "").strip().lower() for c in next(rows, ())]
    if "keyword" not in header:
        raise ValueError(f"{path.name}: the Keywords sheet needs a 'Keyword' column")
    kw_col = header.index("keyword")
    mp_col = header.index("marketplaces") if "marketplaces" in header else None

    out: list[tuple[str, tuple[str, ...]]] = []
    seen: set[str] = set()
    for row in rows:
        keyword = " ".join(str(row[kw_col] or "").split()) if kw_col < len(row) else ""
        if not keyword or keyword.lower() in seen:
            continue
        seen.add(keyword.lower())
        raw = str(row[mp_col] or "") if mp_col is not None and mp_col < len(row) else ""
        wanted = tuple(m for m in MARKETPLACES if m in raw.lower()) or MARKETPLACES
        out.append((keyword, wanted))
    return out


def search_due(last: datetime | None, now: datetime | None = None) -> bool:
    """Once a day, on the first run from SEARCH_FROM_HOUR. Keyed on the last search actually
    stored rather than on the clock alone, so a PC that was off at 9 searches at 10 instead of
    skipping the day."""
    now = now or datetime.now()
    if now.hour < SEARCH_FROM_HOUR:
        return False
    return last is None or last.astimezone().date() < now.date()


# ── doing it ────────────────────────────────────────────────────────────────────────────


def search(
    keywords: list[tuple[str, tuple[str, ...]]],
    fetcher: Fetcher,
    products: list[Product],
) -> tuple[list[SearchResult], list[str]]:
    """Run every keyword on its marketplaces. Returns (results, problems).

    One thread per marketplace, keywords in sequence behind the same per-host throttle the
    price scrape uses. A keyword that fails is reported and skipped — it never fails the run.
    """

    def one_marketplace(marketplace: str) -> tuple[list[SearchResult], list[str]]:
        found: list[SearchResult] = []
        problems: list[str] = []
        for keyword, where in keywords:
            if marketplace not in where:
                continue
            url, body = search_request(marketplace, keyword)
            result = fetcher.fetch_html(marketplace, url, json_body=body)
            if not result.ok or result.html is None:
                problems.append(f"{marketplace} '{keyword}': {result.note or 'no response'}")
                continue
            try:
                page = PARSERS[marketplace](result.html)
            except Exception as exc:  # a changed page must not cost the other keywords
                problems.append(f"{marketplace} '{keyword}': {type(exc).__name__}: {exc}")
                continue
            if not page:
                problems.append(f"{marketplace} '{keyword}': no results read from the page")
            for r in page:
                r.keyword = keyword
            found.extend(page)
        return found, problems

    results: list[SearchResult] = []
    problems: list[str] = []
    with ThreadPoolExecutor(max_workers=len(MARKETPLACES)) as pool:
        for found, issues in pool.map(one_marketplace, MARKETPLACES):
            results.extend(found)
            problems.extend(issues)

    attribute(results, products)
    return results, problems


def _main(argv: list[str]) -> int:
    from .catalogue import load_products
    from .config import load_settings

    if sys.stdout is not None:  # None under pythonw.exe — see run.main()
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    keywords = [(k, MARKETPLACES) for k in argv] or load_keywords()
    products = load_products()
    with Fetcher(load_settings()) as fetcher:
        results, problems = search(keywords, fetcher, products)

    for keyword, _ in keywords:
        for marketplace in MARKETPLACES:
            rows = [r for r in results if r.keyword == keyword and r.marketplace == marketplace]
            if not rows:
                continue
            ours = [r for r in rows if r.brand == "IMECO"]
            best = min((r.position for r in ours), default=None)
            print(
                f"{keyword:<22} {marketplace:<9} {len(rows):>3} results  "
                f"IMECO: {len(ours)}  best #{best if best else '-'}"
            )
    for p in problems:
        print("  !", p)
    return 0


if __name__ == "__main__":
    raise SystemExit(_main(sys.argv[1:]))
