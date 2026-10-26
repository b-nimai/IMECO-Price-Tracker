"""Reads `Docs/pricing link.xlsx` into normalised Products.

The workbook is hand-maintained, so this file is mostly about absorbing its inconsistencies
without letting any of them become a wrong number downstream. Every quirk handled here was
observed in the real file, not guessed at.

Run directly to inspect what was parsed:

    python -m desktop.catalogue
"""

from __future__ import annotations

import os
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from .models import Product

#: The workbook that ships with the app — the repo copy when running from source, and the one
#: baked into the bundle when frozen (`desktop_entry.py` repoints this).
WORKBOOK = Path(__file__).resolve().parent.parent / "Docs" / "pricing link.xlsx"

#: Where an operator-supplied workbook is kept once accepted.
#:
#: Deliberately OUTSIDE the app: a .exe is replaced wholesale on every rebuild, and a sheet
#: stored beside it would be destroyed by the next update. %LOCALAPPDATA% survives that, and
#: survives the exe being moved to another folder.
USER_WORKBOOK = (
    Path(os.environ.get("LOCALAPPDATA", Path.home() / "AppData" / "Local"))
    / "IMECO Price Tracker"
    / "pricing link.xlsx"
)


def active_workbook() -> Path:
    """The workbook a run should read: the operator's if they have supplied one, else ours.

    This is the whole "update the spreadsheet" feature in one function. Nothing else needs to
    know which copy won — `load_products()` calls this, so a scheduled run picks up a sheet the
    operator swapped in an hour ago without the app being rebuilt or even restarted.
    """
    return USER_WORKBOOK if USER_WORKBOOK.exists() else WORKBOOK


def workbook_status() -> tuple[Path, bool, str]:
    """(path, is_operator_supplied, human date) — for the GUI's status line.

    A date is reported only for an OPERATOR-SUPPLIED sheet, because only that one has a
    meaningful date. PyInstaller extracts the bundled workbook into a fresh temporary
    directory on every launch, so its mtime is the moment the app was opened — the built-in
    sheet was displaying "07 Aug 2026 19:46" when 19:46 was simply when someone double-clicked
    the .exe. A timestamp that changes every launch is worse than none: it reads as "this sheet
    is up to date" no matter how old the bundled data actually is.
    """
    path = active_workbook()
    if path != USER_WORKBOOK:
        return path, False, "shipped with this app"
    try:
        stamp = "uploaded " + datetime.fromtimestamp(path.stat().st_mtime).strftime("%d %b %Y %H:%M")
    except OSError:
        stamp = "uploaded, date unknown"
    return path, True, stamp


def install_workbook(source: Path) -> list[Product]:
    """Validate `source`, then adopt it as the workbook every future run will read.

    PARSED BEFORE IT IS COPIED, and that ordering is the point. This app runs unattended once
    an hour; if a malformed sheet could be installed, the next 24 scrapes would fail silently
    while the operator believed they had just updated prices. A workbook that cannot be read,
    or that yields no products, is rejected and the previous one stays in force.
    """
    products = load_products(source)
    if not products:
        raise ValueError("that sheet parsed but contained no products — nothing was changed")

    USER_WORKBOOK.parent.mkdir(parents=True, exist_ok=True)
    # Copy to a temporary name and rename, so a failure part-way through cannot leave a
    # truncated workbook in place of the working one.
    staging = USER_WORKBOOK.with_suffix(".incoming")
    shutil.copyfile(source, staging)
    staging.replace(USER_WORKBOOK)
    return products


def revert_workbook() -> bool:
    """Drop the operator's copy and fall back to the one shipped with the app."""
    if not USER_WORKBOOK.exists():
        return False
    USER_WORKBOOK.unlink()
    return True


def summarise_catalogue(products: list[Product]) -> str:
    """A short "what this sheet contains" line, shown before a swap is committed."""
    counts: dict[str, int] = {}
    for p in products:
        for marketplace in p.urls:
            counts[marketplace] = counts.get(marketplace, 0) + 1
    per = ", ".join(f"{m} {n}" for m, n in sorted(counts.items(), key=lambda kv: -kv[1]))
    own = sum(p.competitor_of is None for p in products)
    rivals = len(products) - own
    who = f"{own} products" + (f" + {rivals} competitor{'s' * (rivals != 1)}" if rivals else "")
    return f"{who}, {sum(len(p.urls) for p in products)} listings ({per})"

# Column A carries a U+2605 star on 24 of 32 rows. It has to come off before the SKU can be
# used as a database key, but it is kept as `Product.starred`.
STAR = "★"

# Header text -> (marketplace, kind). Matched case-insensitively on the stripped header, so
# adding "Blinkit SP" / "Blinkit Link" to the sheet is all that is needed to light Blinkit up.
HEADERS: dict[str, tuple[str, str]] = {
    "amazon sp": ("amazon", "sp"),
    "amazon link": ("amazon", "link"),
    "flipkart sp": ("flipkart", "sp"),
    "flipkart link": ("flipkart", "link"),
    # The sheet still has "Myntra SP" / "Myntra Link" columns. They are deliberately NOT mapped:
    # Myntra was dropped, and an unmapped header is simply ignored, so the links never reach
    # `Product.urls` and never get counted as listings to fetch.
    "blinkit sp": ("blinkit", "sp"),
    "blinkit link": ("blinkit", "link"),
}


def parse_price(raw: object) -> float | None:
    """Cell -> number, or None for a genuinely absent value.

    None rather than 0.0 is deliberate and load-bearing: a blank price cell means "not listed
    here", which must stay distinguishable from a real price of zero all the way to the screen.
    """
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        return float(raw) if raw == raw else None  # NaN check
    s = str(raw).strip()
    if not s or s.lower() in {"nan", "-", "na", "n/a"}:
        return None
    s = re.sub(r"[^\d.]", "", s.replace(",", ""))
    if not s or s == ".":
        return None
    try:
        return float(s)
    except ValueError:
        return None


# ── URL canonicalisation ────────────────────────────────────────────────────────────────
# Each marketplace gets reduced to its stable identifier and rebuilt. This strips tracking
# junk (one Amazon row is a 900-char search URL) and guarantees a consistent shape to fetch.

_CANONICALISERS: dict[str, tuple[re.Pattern[str], str]] = {
    # ASIN. Handles both /dp/B0XXXXXXXX and the /IMECO-Baby-Wipes/dp/B0XXXXXXXX/ref=... form.
    "amazon": (re.compile(r"/dp/([A-Z0-9]{10})"), "https://www.amazon.in/dp/{}"),
    # Blinkit product id.
    "blinkit": (re.compile(r"/prid/(\d+)"), "https://blinkit.com/prn/-/prid/{}"),
}


def canonical_url(marketplace: str, raw: str) -> str | None:
    """Reduce a sheet URL to a canonical, fetchable form. None if it is unusable."""
    url = raw.strip()
    if not url.lower().startswith(("http://", "https://")):
        return None

    if marketplace in _CANONICALISERS:
        pattern, template = _CANONICALISERS[marketplace]
        m = pattern.search(url)
        return template.format(m.group(1)) if m else url

    if marketplace == "flipkart":
        # The real identifier is the pid. All but one row uses the placeholder
        # `/product/p/itme?pid=` slug, which Flipkart resolves fine, so normalise onto it.
        m = re.search(r"[?&]pid=([A-Z0-9]+)", url)
        return f"https://www.flipkart.com/product/p/itme?pid={m.group(1)}" if m else url

    return url


def load_products(path: Path | None = None) -> list[Product]:
    """Parse the workbook. Raises if the file or its header row is missing.

    Defaults to `active_workbook()`, NOT to the shipped copy — so an operator-supplied sheet is
    what every run reads, without a single call site having to know that.
    """
    path = path or active_workbook()
    if not path.exists():
        raise FileNotFoundError(f"catalogue workbook not found: {path}")

    book = load_workbook(path, data_only=True, read_only=True)
    products = _own_products(book.worksheets[0], path)
    rivals = next((ws for ws in book.worksheets if ws.title.strip().lower() == COMPETITOR_SHEET), None)
    if rivals is not None:
        products += _competitors(rivals, path, products)

    _reject_duplicates(products, path)
    return products


def _header(rows, path: Path, sheet: str) -> tuple[dict[str, int], dict[int, tuple[str, str]]]:
    """(fixed column -> index, index -> (marketplace, kind)) from a sheet's first row.

    Driven by header text rather than position so inserting a column does not silently shift
    every price.
    """
    try:
        header = next(rows)
    except StopIteration as exc:
        raise ValueError(f"{path.name}: the {sheet} sheet is empty") from exc

    fixed: dict[str, int] = {}
    cols: dict[int, tuple[str, str]] = {}
    for i, cell in enumerate(header):
        key = str(cell or "").strip().lower()
        if key in HEADERS:
            cols[i] = HEADERS[key]
        elif key:
            fixed.setdefault(key, i)
    return fixed, cols


def _read_listings(product: Product, row: tuple, cols: dict[int, tuple[str, str]]) -> None:
    """Fill a product's links and target prices from its row."""
    for i, (marketplace, kind) in cols.items():
        if i >= len(row):
            continue
        value = row[i]
        if value is None:
            continue
        if kind == "sp":
            price = parse_price(value)
            if price is not None:
                product.target_sp[marketplace] = price
        else:
            url = canonical_url(marketplace, str(value))
            if url:
                product.urls[marketplace] = url


def _cell(row: tuple, index: int | None) -> object:
    return row[index] if index is not None and index < len(row) else None


def _own_products(ws, path: Path) -> list[Product]:
    """IMECO's own catalogue: the workbook's first sheet."""
    rows = ws.iter_rows(values_only=True)
    fixed, cols = _header(rows, path, "first")
    sku_col, name_col, mrp_col = fixed.get("sku"), fixed.get("product name"), fixed.get("mrp")
    if sku_col is None or name_col is None:
        raise ValueError(f"{path.name}: missing required 'SKU' / 'Product Name' headers")

    products: list[Product] = []
    for row in rows:
        raw_sku = str(_cell(row, sku_col) or "").replace(STAR, "").strip()
        if not raw_sku:
            continue  # blank spacer row

        product = Product(
            sku=raw_sku,
            name=str(_cell(row, name_col) or "").strip(),
            mrp=parse_price(_cell(row, mrp_col)),
            starred=STAR in str(_cell(row, sku_col) or ""),
            # Position among PARSED rows, not among worksheet rows: blank spacers are skipped
            # above, and a gap in the numbering would otherwise appear wherever the sheet has
            # one. The dashboard only ever sorts on this, so density is what matters, not
            # matching a row number in Excel.
            position=len(products),
        )
        _read_listings(product, row, cols)
        products.append(product)
    return products


# ── Competitors ─────────────────────────────────────────────────────────────────────────
# An optional second sheet, named "Competitors", listing rival products to price against
# IMECO's. One row per competing product:
#
#   Competitor Of | Brand | Product Name | MRP | Amazon Link | Flipkart Link | Blinkit Link
#
# "Competitor Of" is the IMECO SKU it competes with. There is deliberately no SKU column to
# fill in: a competitor's key is made from its first link (`C:amazon:B0XXXXXXXX`), which is
# already unique and stays the same however the row is edited or moved. An explicit "SKU"
# column is still honoured if someone adds one.
COMPETITOR_SHEET = "competitors"
_LISTING_ID = {
    "amazon": re.compile(r"/dp/([A-Z0-9]{10})"),
    "flipkart": re.compile(r"[?&]pid=([A-Z0-9]+)"),
    "blinkit": re.compile(r"/prid/(\d+)"),
}


def _competitor_key(product: Product) -> str | None:
    for marketplace in ("amazon", "flipkart", "blinkit"):
        url = product.urls.get(marketplace)
        m = url and _LISTING_ID[marketplace].search(url)
        if m:
            return f"C:{marketplace}:{m.group(1)}"
    return None


def _competitors(ws, path: Path, own: list[Product]) -> list[Product]:
    """The "Competitors" sheet. Refused, by row, if a row points at an IMECO SKU that does not
    exist — the same reasoning as `_reject_duplicates`: tell the operator now, while they are
    looking at the file, not an hour later inside a scheduled run."""
    rows = ws.iter_rows(values_only=True)
    fixed, cols = _header(rows, path, "Competitors")
    of_col = fixed.get("competitor of")
    if of_col is None:
        raise ValueError(f"{path.name}: the Competitors sheet needs a 'Competitor Of' column")
    brand_col, name_col = fixed.get("brand"), fixed.get("product name")
    mrp_col, sku_col = fixed.get("mrp"), fixed.get("sku")

    known = {p.sku for p in own}
    out: list[Product] = []
    for n, row in enumerate(rows, start=2):  # Excel row number, header being row 1
        of = str(_cell(row, of_col) or "").replace(STAR, "").strip()
        if not of:
            continue  # blank spacer row
        if of not in known:
            raise ValueError(
                f"{path.name}: Competitors row {n} says 'Competitor Of' {of!r}, which is not a "
                "SKU on the first sheet. Use the IMECO SKU exactly as it is written there."
            )

        product = Product(
            sku="",
            name=str(_cell(row, name_col) or "").strip(),
            mrp=parse_price(_cell(row, mrp_col)),
            starred=False,
            # After every IMECO product, so a competitor can never sit among them in sheet order.
            position=len(own) + len(out),
            brand=str(_cell(row, brand_col) or "").strip() or "Unknown brand",
            competitor_of=of,
        )
        _read_listings(product, row, cols)

        explicit = str(_cell(row, sku_col) or "").strip()
        key = explicit or _competitor_key(product)
        if not key:
            raise ValueError(
                f"{path.name}: Competitors row {n} has no Amazon, Flipkart or Blinkit link."
            )
        product.sku = key
        if not product.name:
            product.name = f"{product.brand} ({key})"
        out.append(product)
    return out


def _reject_duplicates(products: list[Product], path: Path) -> None:
    """A SKU may appear once. Refuse the sheet, by name, if one appears twice.

    The SKU is the database's primary key, so a repeat makes the catalogue upsert fail with
    `ON CONFLICT DO UPDATE command cannot affect row a second time` — a message that says
    nothing to the person who typed the duplicate, arriving an hour later inside a scheduled
    run that then stores nothing at all.

    Raised from the parser so `install_workbook` catches it BEFORE the sheet is adopted: the
    operator is told which SKU to fix while they are still looking at the file picker, and the
    previous workbook stays in force meanwhile.
    """
    seen: set[str] = set()
    repeated: list[str] = []
    for product in products:
        if product.sku in seen and product.sku not in repeated:
            repeated.append(product.sku)
        seen.add(product.sku)

    if repeated:
        listed = ", ".join(repeated)
        if all(s.startswith("C:") for s in repeated):
            # Competitor keys are made from the link, so a repeat means the same listing twice.
            raise ValueError(
                f"{path.name}: the Competitors sheet lists the same product link twice ({listed}). "
                "Keep one row per competitor product and try again."
            )
        raise ValueError(
            f"{path.name}: the same SKU appears on more than one row ({listed}). "
            "Every product needs its own SKU — merge or rename those rows and try again."
        )


def _main() -> int:
    """Inspect the parsed catalogue — the check described in the plan's Verification step 1."""
    if sys.stdout is not None:  # None under pythonw.exe — see run.main()
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    products = load_products()

    print(f"{len(products)} products from {WORKBOOK.name}\n")
    counts: dict[str, int] = {}
    for p in products:
        for marketplace in p.urls:
            counts[marketplace] = counts.get(marketplace, 0) + 1

    print("links per marketplace:")
    for marketplace, n in sorted(counts.items(), key=lambda kv: -kv[1]):
        print(f"  {marketplace:<9} {n}")
    print(f"\nstarred: {sum(p.starred for p in products)}")
    print(f"competitors: {sum(p.competitor_of is not None for p in products)}")
    print(f"total listings to fetch: {sum(len(p.urls) for p in products)}\n")

    for p in products[:3]:
        print(f"  {p.sku}  mrp={p.mrp}  starred={p.starred}")
        print(f"    {p.name[:70]}")
        for marketplace, url in p.urls.items():
            sp = p.target_sp.get(marketplace)
            print(f"      {marketplace:<9} sp={sp!s:<7} {url[:78]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(_main())
